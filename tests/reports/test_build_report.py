"""Testes do gerador de relatório, com resultados sintéticos dos dois frameworks.

Herméticos: não leem tests/*/results nem chamam Ollama. Rodar da raiz:
    tests\\reports\\.venv\\Scripts\\python.exe -m pytest tests/reports
"""
from __future__ import annotations

import csv
import json
from pathlib import Path

import pytest
from openpyxl import load_workbook

import build_report
import report_charts
import report_data
from report_data import NOT_FOUND_MARKER, compare, dimension_stats, fmt, load_run, signed
from report_document import Para
from report_markdown import render_markdown

DEEPEVAL_NAMES = {
    "faithfulness": "Faithfulness",
    "answer_relevancy": "Answer Relevancy",
    "context_precision": "Contextual Precision",
    "context_recall": "Contextual Recall",
    "context_relevance": "Contextual Relevancy",
}


def _metrics(scores: dict[str, float | None], deepeval: bool) -> list[dict]:
    out = []
    for key, value in scores.items():
        out.append(
            {
                "name": DEEPEVAL_NAMES[key] if deepeval else key,
                "score": value,
                "success": None if value is None else value >= 0.7,
                "reason": "porque sim" if deepeval and value is not None else None,
                "error": "timeout: sem resposta em 300s" if value is None else None,
            }
        )
    return out


def _case(name: str, answer: str, scores: dict[str, float | None], deepeval: bool, contexts=("t1", "t2")) -> dict:
    return {
        "name": name,
        "input": f"Pergunta {name}?",
        "actual_output": answer,
        "expected_output": f"Referência {name}.",
        "retrieval_context": list(contexts),
        "success": all(v is not None and v >= 0.7 for v in scores.values()),
        "metrics": _metrics(scores, deepeval),
    }


ALL_HIGH = dict.fromkeys(DEEPEVAL_NAMES, 0.9)


@pytest.fixture
def ragas_file(tmp_path: Path) -> Path:
    data = {
        "framework": "ragas",
        "ragas_version": "0.4.3",
        "generated_at": "2026-10-07T11:51:07+00:00",
        "generation_model": "gemma3:4b",
        "judge_model": "qwen3:8b",
        "embedding_model": "nomic-embed-text",
        "threshold": 0.7,
        "prompt_overrides": {"answer_relevancy": "instruction += 'x'"},
        "judge_calls": 10,
        "judge_failures": 0,
        "n_cases": 2,
        "cases": [
            _case("c1", "Resposta um.", ALL_HIGH, deepeval=False),
            _case("c2", f"{NOT_FOUND_MARKER}.", {**ALL_HIGH, "answer_relevancy": 0.0}, deepeval=False),
        ],
    }
    path = tmp_path / "20261007T115107Z.json"
    path.write_text(json.dumps(data), encoding="utf-8")
    return path


@pytest.fixture
def deepeval_file(tmp_path: Path) -> Path:
    data = {
        "generated_at": "2026-08-27T17:27:33+00:00",
        "generation_model": "gemma3:4b",
        "judge_model": "gemma3:4b",
        "threshold": 0.7,
        "n_cases": 2,
        "cases": [
            _case("c1", "Resposta um.", {**ALL_HIGH, "faithfulness": 0.2}, deepeval=True),
            _case("c2", "Outra resposta.", {**ALL_HIGH, "context_recall": None}, deepeval=True, contexts=("t9",)),
        ],
    }
    path = tmp_path / "20260827T172733Z.json"
    path.write_text(json.dumps(data), encoding="utf-8")
    return path


def _jev_metric(name: str, value: float | None, applicable: bool = True) -> dict:
    return {
        "name": name,
        "score": value,
        "success": None if value is None else value >= (0.5 if name == "correctness" else 0.7),
        "applicable": applicable,
        "reason": "Decided by jev-1.13.0 (TypeSafe AI)." if applicable else "N/A: nenhuma pergunta se aplica.",
        "error": None,
        "confidence": 0.8,
    }


@pytest.fixture
def jeveval_file(tmp_path: Path) -> Path:
    # O formato de tests/jeveval/run_and_export.py: c2 é uma recusa, e fidelidade
    # e relevância da resposta saem N/A (não 0, não n/d).
    contextos = {"context_precision": 0.9, "context_recall": 0.9, "context_relevance": 0.9}
    data = {
        "framework": "jeveval",
        "deepeval_version": "4.3.1",
        "generated_at": "2026-10-16T12:00:00+00:00",
        "generation_model": "gemma3:4b",
        "judge_model": "jev-1.13.0 (TypeSafe AI)",
        "threshold": 0.7,
        "instrument_sha": "abc123def456",
        "jev_calls": 12,
        "jev_failures": 0,
        "cost_usd": 0.0012,
        "n_cases": 2,
        "cases": [
            {
                **_case("c1", "Resposta um.", {}, deepeval=False),
                "metrics": [_jev_metric("faithfulness", 0.9), _jev_metric("answer_relevancy", 0.9)]
                + [_jev_metric(k, v) for k, v in contextos.items()]
                + [_jev_metric("correctness", 0.95)],
            },
            {
                **_case("c2", f"{NOT_FOUND_MARKER}.", {}, deepeval=False),
                "metrics": [_jev_metric("faithfulness", None, False), _jev_metric("answer_relevancy", None, False)]
                + [_jev_metric(k, v) for k, v in contextos.items()]
                + [_jev_metric("correctness", 0.1)],
            },
        ],
    }
    path = tmp_path / "20261016T120000Z.json"
    path.write_text(json.dumps(data), encoding="utf-8")
    return path


def test_marcador_de_abstencao_vem_do_backend() -> None:
    # Lido de backend/app/rag/generation.py por ast; se o nome ou o formato
    # mudar lá, a detecção de abstenção para em silêncio sem este teste.
    assert NOT_FOUND_MARKER is not None
    assert "não encontrei" in NOT_FOUND_MARKER.lower()


def test_detecta_o_framework_de_cada_formato(
    ragas_file: Path, deepeval_file: Path, jeveval_file: Path, tmp_path: Path
) -> None:
    assert load_run(ragas_file).framework == "RAGAS"
    assert load_run(deepeval_file).framework == "DeepEval"
    assert load_run(jeveval_file).framework == "JevEval"

    estranho = tmp_path / "x.json"
    estranho.write_text(json.dumps({"cases": [{"metrics": [{"name": "BLEU"}]}]}), encoding="utf-8")
    with pytest.raises(ValueError):
        load_run(estranho)


def test_mapeia_as_metricas_do_deepeval_para_as_dimensoes(deepeval_file: Path) -> None:
    run = load_run(deepeval_file)

    c1, c2 = run.cases
    assert c1.score("faithfulness").value == pytest.approx(0.2)
    assert c1.score("faithfulness").passed is False
    assert c1.score("faithfulness").reason == "porque sim"
    assert not c2.score("context_recall").measured
    assert "timeout" in c2.score("context_recall").error


def test_nd_fica_fora_da_media_mas_nao_conta_como_aprovada(deepeval_file: Path) -> None:
    stats = dimension_stats(load_run(deepeval_file))["context_recall"]

    assert stats.mean == pytest.approx(0.9)  # só o c1 foi medido
    assert (stats.n_measured, stats.n_passed, stats.n_cases) == (1, 1, 2)
    assert stats.errors == [("c2", "timeout: sem resposta em 300s")]


def test_na_do_jeveval_sai_das_medias_dos_denominadores_e_dos_erros(jeveval_file: Path) -> None:
    run = load_run(jeveval_file)
    stats = dimension_stats(run)["faithfulness"]

    assert (stats.n_cases, stats.n_measured, stats.n_passed, stats.n_not_applicable) == (1, 1, 1, 1)
    assert stats.mean == pytest.approx(0.9)
    assert stats.errors == []  # N/A não é falha de medida
    c2 = run.case("c2")
    assert not c2.score("faithfulness").applicable
    assert c2.all_passed  # as dimensões que se aplicam passaram; N/A não reprova


def test_nd_continua_reprovando_quando_nao_e_na(deepeval_file: Path) -> None:
    # O N/A não pode afrouxar a regra antiga: nota não medida segue sem aprovar.
    c2 = load_run(deepeval_file).case("c2")

    assert c2.score("context_recall").applicable
    assert not c2.all_passed


def test_correcao_do_jeveval_fica_fora_das_cinco_dimensoes(jeveval_file: Path) -> None:
    c1 = load_run(jeveval_file).case("c1")

    assert set(c1.scores) == {d.key for d in report_data.DIMENSIONS}


def test_detecta_abstencao(ragas_file: Path) -> None:
    c1, c2 = load_run(ragas_file).cases

    assert not c1.abstained
    assert c2.abstained


def test_compara_veredito_por_dimensao_e_avisa_o_que_impede_comparar(ragas_file: Path, deepeval_file: Path) -> None:
    comp = compare(load_run(deepeval_file), load_run(ragas_file))

    faith = comp.per_dimension["faithfulness"]
    assert (faith.n_both, faith.n_agree) == (2, 1)  # c1: 0,2 reprova x 0,9 passa
    assert comp.per_dimension["context_recall"].n_both == 1  # c2 sem nota no DeepEval
    # Só vereditos opostos entram, da maior diferença para a menor:
    # c2/relevância 0,9 -> 0,0 (|0,9|) antes de c1/fidelidade 0,2 -> 0,9 (|0,7|).
    assert [(d.case, d.dimension) for d in comp.divergences] == [("c2", "answer_relevancy"), ("c1", "faithfulness")]
    assert comp.divergences[1].diff == pytest.approx(0.7)
    assert comp.identical_answers == 1
    assert comp.identical_contexts == 1

    avisos = " ".join(comp.warnings)
    assert "Juízes diferentes" in avisos
    assert "Só 1 de 2 respostas" in avisos


def test_cor_segue_o_framework_e_nao_a_ordem(ragas_file: Path, deepeval_file: Path, jeveval_file: Path) -> None:
    ragas, deepeval, jeveval = load_run(ragas_file), load_run(deepeval_file), load_run(jeveval_file)

    assert report_charts.run_colors([ragas, deepeval]) == ["#eb6834", "#2a78d6"]
    assert report_charts.run_colors([deepeval, ragas]) == ["#2a78d6", "#eb6834"]
    assert report_charts.run_colors([jeveval, ragas]) == ["#1baf7a", "#eb6834"]
    # Segunda rodada do mesmo framework: outra cor, fora dos slots reservados
    # (azul, laranja e verde são de DeepEval, RAGAS e JevEval).
    assert report_charts.run_colors([ragas, ragas]) == ["#eb6834", "#eda100"]
    with pytest.raises(ValueError):
        report_charts.run_colors([ragas] * 9)


def test_numeros_no_formato_do_texto() -> None:
    assert fmt(0.5) == "0,50"
    assert fmt(None) == "—"
    assert signed(-0.5) == "−0,50"
    assert signed(0.25) == "+0,25"


def test_markdown_escapa_tag_vinda_da_resposta(tmp_path: Path) -> None:
    path = render_markdown([Para("O modelo escreveu <think> no meio.")], tmp_path / "r.md")

    assert "&lt;think>" in path.read_text(encoding="utf-8")


def test_gera_os_quatro_formatos(ragas_file: Path, deepeval_file: Path, tmp_path: Path) -> None:
    out = tmp_path / "out"
    outputs = build_report.build([ragas_file, deepeval_file], list(build_report.FORMATS), out)

    assert set(outputs) == {"md", "pdf", "xlsx", "csv"}
    assert outputs["pdf"].read_bytes()[:5] == b"%PDF-"

    md = outputs["md"].read_text(encoding="utf-8")
    assert "DeepEval (27/08/2026" in md
    assert "RAGAS (07/10/2026" in md
    assert "Comparabilidade" in md
    assert "figuras/medias_por_dimensao.png" in md
    for figure in (out / "figuras").iterdir():
        assert figure.stat().st_size > 1000

    wb = load_workbook(outputs["xlsx"])
    assert {"Resumo", "Médias", "Notas", "Matriz DeepEval", "Matriz RAGAS", "Concordância", "Divergências",
            "Respostas"} <= set(wb.sheetnames)
    assert wb["Notas"].max_row == 1 + 2 * 2 * 5

    with outputs["csv"].open(encoding="utf-8-sig") as f:
        rows = list(csv.reader(f, delimiter=";"))
    assert len(rows) == 1 + 2 * 2 * 5


def test_relatorio_com_os_tres_frameworks(
    ragas_file: Path, deepeval_file: Path, jeveval_file: Path, tmp_path: Path
) -> None:
    out = tmp_path / "out"
    outputs = build_report.build([jeveval_file, ragas_file, deepeval_file], list(build_report.FORMATS), out)

    md = outputs["md"].read_text(encoding="utf-8")
    assert "JevEval (16/10/2026" in md
    assert md.index("DeepEval (27/08/2026") < md.index("RAGAS (07/10/2026") < md.index("JevEval (16/10/2026")
    assert "N/A" in md
    assert "`JevEval(faithfulness)`" in md
    assert "Instrumento JevEval `abc123def456`" in md
    assert "Chamadas ao Jev: 12" in md
    assert "nota(s) N/A" in md
    assert "Concordância: RAGAS × JevEval" in md
    assert "Matriz JevEval" in load_workbook(outputs["xlsx"]).sheetnames
    with outputs["csv"].open(encoding="utf-8-sig") as f:
        verdicts = [row[8] for row in csv.reader(f, delimiter=";")]
    assert verdicts.count("N/A") == 2


def test_uma_rodada_so_gera_relatorio_sem_comparacao(ragas_file: Path, tmp_path: Path) -> None:
    outputs = build_report.build([ragas_file], ["md", "xlsx"], tmp_path / "out")

    md = outputs["md"].read_text(encoding="utf-8")
    assert "Comparação entre frameworks" not in md
    assert "Concordância" not in load_workbook(outputs["xlsx"]).sheetnames


def test_descobre_o_resultado_mais_recente(tmp_path: Path) -> None:
    pasta = tmp_path / "results"
    pasta.mkdir()
    for nome in ("20260101T000000Z.json", "20261007T115107Z.json", "latest.html", "rascunho.json"):
        (pasta / nome).write_text("{}", encoding="utf-8")

    assert report_data.latest_result(pasta).name == "20261007T115107Z.json"
