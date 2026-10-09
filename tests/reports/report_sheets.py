"""Planilha (.xlsx) e CSV do relatório — os dados, em formato de análise.

Diferente do Markdown/PDF (texto para ler), aqui vão tabelas para filtrar,
pivotar e colar em outra análise: notas em formato longo (uma linha por
rodada x caso x dimensão), médias, matrizes caso x dimensão com a mesma
escala divergente dos gráficos (formatação condicional nativa, então a cor
acompanha se alguém editar a nota), concordância e divergências entre
rodadas, e as respostas completas.
"""
from __future__ import annotations

import csv
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from report_data import (
    DIMENSION_BY_KEY,
    DIMENSIONS,
    NATIVE_METRIC,
    Comparison,
    Run,
    dimension_stats,
    run_highlights,
)

HEADER_FILL = PatternFill("solid", fgColor="F0EFEC")
HEADER_FONT = Font(bold=True, color="0B0B0B")
HEADER_BORDER = Border(bottom=Side(style="thin", color="C3C2B7"))
TITLE_FONT = Font(bold=True, size=14)
WRAP = Alignment(wrap_text=True, vertical="top")
SCORE_FORMAT = "0.00"

LONG_COLUMNS = [
    "Rodada",
    "Framework",
    "Data",
    "Caso",
    "Dimensão",
    "Métrica",
    "Nota",
    "Veredito",
    "Abstenção",
    "Erro",
    "Justificativa do juiz",
]


def _verdict(passed: bool | None, measured: bool) -> str:
    if not measured:
        return "n/d"
    return "passou" if passed else "reprovou"


def long_rows(runs: list[Run]) -> list[list]:
    """Uma linha por rodada x caso x dimensão (a tabela "Notas" e o CSV)."""
    rows = []
    for run in runs:
        for case in run.cases:
            for dim in DIMENSIONS:
                s = case.score(dim.key)
                rows.append(
                    [
                        run.label,
                        run.framework,
                        run.when,
                        case.name,
                        dim.label,
                        NATIVE_METRIC.get(run.framework, {}).get(dim.key, dim.key),
                        None if s.value is None else round(s.value, 4),
                        _verdict(s.passed, s.measured),
                        "sim" if case.abstained else "não",
                        s.error or "",
                        s.reason or "",
                    ]
                )
    return rows


def write_csv(runs: list[Run], path: Path) -> Path:
    # ';' e BOM: abre direto no Excel em português. Nota com ponto decimal, para
    # pandas/R lerem sem configuração (na planilha .xlsx ela já é número).
    with path.open("w", newline="", encoding="utf-8-sig") as f:
        writer = csv.writer(f, delimiter=";")
        writer.writerow(LONG_COLUMNS)
        writer.writerows(long_rows(runs))
    return path


# --------------------------------------------------------------------- xlsx


def _header(ws: Worksheet, row: int, headers: list[str], col: int = 1) -> None:
    for i, text in enumerate(headers):
        cell = ws.cell(row=row, column=col + i, value=text)
        cell.fill, cell.font, cell.border = HEADER_FILL, HEADER_FONT, HEADER_BORDER
        cell.alignment = Alignment(vertical="center", wrap_text=True)


def _widths(ws: Worksheet, widths: list[float]) -> None:
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w


def _diverging(ws: Worksheet, ref: str, threshold: float) -> None:
    """Mesma escala dos gráficos: vermelho em 0, cinza no limiar, azul em 1."""
    ws.conditional_formatting.add(
        ref,
        ColorScaleRule(
            start_type="num", start_value=0, start_color="E34948",
            mid_type="num", mid_value=threshold, mid_color="F0EFEC",
            end_type="num", end_value=1, end_color="2A78D6",
        ),
    )


def _table(ws: Worksheet, headers: list[str], rows: list[list], widths: list[float], start_row: int = 1) -> int:
    _header(ws, start_row, headers)
    for r, row in enumerate(rows, start=start_row + 1):
        for c, value in enumerate(row, start=1):
            cell = ws.cell(row=r, column=c, value=value)
            if isinstance(value, float):
                cell.number_format = SCORE_FORMAT
    _widths(ws, widths)
    return start_row + len(rows)


def _sheet_name(base: str, taken: set[str]) -> str:
    name, n = base[:31], 2
    while name in taken:
        suffix = f" ({n})"
        name, n = base[: 31 - len(suffix)] + suffix, n + 1
    taken.add(name)
    return name


def _execution_sheet(wb: Workbook, taken: set[str], execution: dict) -> None:
    ws = wb.create_sheet(_sheet_name("Execução", taken), index=1)
    rows = []
    for step in execution.get("steps", []):
        tests = step.get("tests") or {}
        rows.append(
            [
                step["name"],
                step["status"],
                tests.get("total"),
                (tests.get("failed") or 0) + (tests.get("errors") or 0) if tests else None,
                tests.get("skipped"),
                step.get("seconds"),
                step.get("detail") or "",
                step.get("log") or "",
            ]
        )
    _table(ws, ["Suíte", "Resultado", "Testes", "Falhas", "Pulados", "Duração (s)", "Observação", "Log"], rows,
           [30, 11, 9, 9, 9, 12, 70, 28])
    ws.freeze_panes = "A2"


def _history_sheet(wb: Workbook, taken: set[str], history: list[dict], threshold: float) -> None:
    ws = wb.create_sheet(_sheet_name("Histórico", taken))
    rows = []
    for run in reversed(history):
        rows.append(
            [
                run.get("run_generated_at"),
                run["framework"],
                run.get("generation_model"),
                run.get("judge_model"),
                *(None if run["means"].get(d.key) is None else round(run["means"][d.key], 4) for d in DIMENSIONS),
                run.get("all_passed"),
                run.get("n_cases"),
                "sim" if run.get("shared_answers") else "não",
            ]
        )
    end = _table(
        ws,
        ["Data da avaliação", "Framework", "Geração", "Juiz"] + [d.label for d in DIMENSIONS]
        + ["Casos com todas ≥ limiar", "Casos", "Respostas compartilhadas"],
        rows,
        [26, 11, 14, 12] + [14] * len(DIMENSIONS) + [14, 8, 14],
    )
    last = get_column_letter(4 + len(DIMENSIONS))
    _diverging(ws, f"E2:{last}{end}", threshold)
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = f"A1:{get_column_letter(4 + len(DIMENSIONS) + 3)}{end}"


def write_xlsx(
    runs: list[Run],
    comparisons: list[Comparison],
    figures: dict[str, Path],
    generated_at: datetime,
    path: Path,
    execution: dict | None = None,
    history: list[dict] | None = None,
) -> Path:
    wb = Workbook()
    taken: set[str] = set()

    # Resumo -----------------------------------------------------------------
    ws = wb.active
    ws.title = _sheet_name("Resumo", taken)
    ws["A1"] = "Relatório de qualidade de resposta do RAG"
    ws["A1"].font = TITLE_FONT
    ws["A2"] = f"Gerado em {generated_at.strftime('%d/%m/%Y %H:%M')}"
    rows = []
    for run in runs:
        stats = dimension_stats(run)
        rows.append(
            [
                run.framework,
                run.when,
                run.generation_model,
                run.judge_model,
                run.embedding_model or "",
                run.threshold,
                len(run.cases),
                sum(s.n_measured for s in stats.values()),
                sum(1 for c in run.cases if c.all_passed),
                str(run.source),
            ]
        )
    end = _table(
        ws,
        ["Framework", "Data", "Modelo de geração", "Juiz", "Embeddings", "Limiar", "Casos", "Notas medidas",
         "Casos com todas ≥ limiar", "Arquivo"],
        rows,
        [12, 17, 18, 14, 18, 8, 8, 13, 14, 70],
        start_row=4,
    )
    row = end + 2
    for run in runs:
        ws.cell(row=row, column=1, value=run.label).font = Font(bold=True)
        row += 1
        for line in run_highlights(run, dimension_stats(run)):
            ws.cell(row=row, column=1, value=f"• {line}")
            row += 1
        row += 1
    for comp in comparisons:
        if comp.warnings:
            ws.cell(row=row, column=1, value=f"⚠ Comparabilidade: {comp.a.label} × {comp.b.label}").font = Font(bold=True)
            row += 1
            for warning in comp.warnings:
                ws.cell(row=row, column=1, value=f"• {warning}")
                row += 1
            row += 1
    if "means" in figures:
        image = XLImage(str(figures["means"]))
        image.width, image.height = image.width * 0.42, image.height * 0.42
        ws.add_image(image, f"A{row + 1}")

    # Médias -----------------------------------------------------------------
    ws = wb.create_sheet(_sheet_name("Médias", taken))
    rows = []
    for run in runs:
        for dim in DIMENSIONS:
            s = dimension_stats(run)[dim.key]
            rows.append(
                [
                    run.label,
                    dim.label,
                    NATIVE_METRIC.get(run.framework, {}).get(dim.key, dim.key),
                    *(None if v is None else round(v, 4) for v in (s.mean, s.median, s.minimum, s.maximum)),
                    s.n_measured,
                    s.n_passed,
                    s.n_cases,
                ]
            )
    end = _table(
        ws,
        ["Rodada", "Dimensão", "Métrica", "Média", "Mediana", "Mín.", "Máx.", "Medidas", "≥ limiar", "Casos"],
        rows,
        [30, 24, 30, 9, 9, 9, 9, 9, 9, 8],
    )
    _diverging(ws, f"D2:G{end}", runs[0].threshold)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:J{end}"

    # Notas (formato longo) --------------------------------------------------
    ws = wb.create_sheet(_sheet_name("Notas", taken))
    end = _table(ws, LONG_COLUMNS, long_rows(runs), [30, 11, 17, 28, 22, 30, 8, 10, 10, 30, 80])
    _diverging(ws, f"G2:G{end}", runs[0].threshold)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:K{end}"

    # Matriz por rodada ------------------------------------------------------
    for run in runs:
        ws = wb.create_sheet(_sheet_name(f"Matriz {run.framework}", taken))
        ws["A1"] = run.label
        ws["A1"].font = Font(bold=True)
        rows = [
            [case.name]
            + [None if case.score(d.key).value is None else round(case.score(d.key).value, 4) for d in DIMENSIONS]
            + ["sim" if case.all_passed else "não", "sim" if case.abstained else "não"]
            for case in run.cases
        ]
        end = _table(
            ws,
            ["Caso"] + [d.label for d in DIMENSIONS] + ["Todas ≥ limiar", "Abstenção"],
            rows,
            [30] + [14] * len(DIMENSIONS) + [14, 11],
            start_row=3,
        )
        last_col = get_column_letter(1 + len(DIMENSIONS))
        _diverging(ws, f"B4:{last_col}{end}", run.threshold)
        ws.freeze_panes = "B4"

    # Concordância e divergências -------------------------------------------
    if comparisons:
        ws = wb.create_sheet(_sheet_name("Concordância", taken))
        rows = []
        for comp in comparisons:
            for dim in DIMENSIONS:
                a = comp.per_dimension[dim.key]
                rows.append(
                    [
                        f"{comp.a.label} × {comp.b.label}",
                        dim.label,
                        a.n_both,
                        a.n_agree,
                        None if a.agreement is None else round(a.agreement, 4),
                        None if a.mean_abs_diff is None else round(a.mean_abs_diff, 4),
                    ]
                )
        end = _table(
            ws,
            ["Par de rodadas", "Dimensão", "Casos com as duas notas", "Mesmo veredito", "Concordância",
             "Diferença média de nota"],
            rows,
            [55, 24, 14, 14, 13, 14],
        )
        for r in range(2, end + 1):
            ws.cell(row=r, column=5).number_format = "0%"
        ws.freeze_panes = "A2"

        ws = wb.create_sheet(_sheet_name("Divergências", taken))
        rows = [
            [
                f"{comp.a.label} × {comp.b.label}",
                d.case,
                DIMENSION_BY_KEY[d.dimension].label,
                round(d.value_a, 4),
                round(d.value_b, 4),
                round(d.diff, 4),
            ]
            for comp in comparisons
            for d in comp.divergences
        ]
        end = _table(ws, ["Par de rodadas", "Caso", "Dimensão", "Nota A", "Nota B", "Diferença (B − A)"],
                     rows, [55, 28, 24, 9, 9, 15])
        ws.freeze_panes = "A2"
        if rows:
            ws.auto_filter.ref = f"A1:F{end}"

    # Respostas ---------------------------------------------------------------
    ws = wb.create_sheet(_sheet_name("Respostas", taken))
    names: list[str] = []
    for run in runs:
        names += [c.name for c in run.cases if c.name not in names]
    rows = []
    for name in names:
        first = next(c for r in runs if (c := r.case(name)))
        rows.append(
            [name, first.question, first.expected] + [(r.case(name).answer if r.case(name) else "") for r in runs]
        )
    end = _table(ws, ["Caso", "Pergunta", "Referência"] + [f"Resposta — {r.label}" for r in runs], rows,
                 [28, 45, 60] + [70] * len(runs))
    for r in range(2, end + 1):
        for c in range(1, 4 + len(runs)):
            ws.cell(row=r, column=c).alignment = WRAP
    ws.freeze_panes = "B2"

    if execution:
        _execution_sheet(wb, taken, execution)
    if history:
        _history_sheet(wb, taken, history, runs[0].threshold)

    wb.save(path)
    return path
