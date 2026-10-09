"""Testes do orquestrador (tests/run_all.py) e do histórico de relatórios.

Herméticos: nenhuma suíte de verdade roda aqui. Rodar da raiz:
    tests\\reports\\.venv\\Scripts\\python.exe -m pytest tests/reports
"""
from __future__ import annotations

import json
import sys
from pathlib import Path

import pytest
from openpyxl import load_workbook

import build_report
import report_history

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))  # tests/, para importar run_all
import run_all  # noqa: E402

from test_build_report import deepeval_file, ragas_file  # noqa: E402,F401 - fixtures reaproveitadas

PYTEST_JUNIT = """<?xml version="1.0" encoding="utf-8"?>
<testsuites><testsuite name="pytest" errors="0" failures="1" skipped="2" tests="10" time="1.2"/></testsuites>"""
PLAYWRIGHT_JUNIT = """<?xml version="1.0" encoding="UTF-8"?>
<testsuites id="" name="" tests="7" failures="0" skipped="1" errors="0" time="12.3">
  <testsuite name="chat-flow.spec.ts" tests="3" failures="0" skipped="1" errors="0"/>
  <testsuite name="monitor.spec.ts" tests="4" failures="0" skipped="0" errors="0"/>
</testsuites>"""


def test_le_junit_do_pytest(tmp_path: Path) -> None:
    path = tmp_path / "api.xml"
    path.write_text(PYTEST_JUNIT, encoding="utf-8")

    assert run_all.junit_counts(path) == {"total": 10, "failed": 1, "errors": 0, "skipped": 2}


def test_le_junit_do_playwright_somando_as_suites_sem_contar_o_total_duas_vezes(tmp_path: Path) -> None:
    path = tmp_path / "e2e.xml"
    path.write_text(PLAYWRIGHT_JUNIT, encoding="utf-8")

    assert run_all.junit_counts(path) == {"total": 7, "failed": 0, "errors": 0, "skipped": 1}


def test_junit_ausente_ou_quebrado_vira_none(tmp_path: Path) -> None:
    assert run_all.junit_counts(tmp_path / "nao_existe.xml") is None
    quebrado = tmp_path / "x.xml"
    quebrado.write_text("<testsuites", encoding="utf-8")
    assert run_all.junit_counts(quebrado) is None


@pytest.mark.parametrize(
    ("code", "tests", "status"),
    [
        (0, {"total": 5, "failed": 0, "errors": 0, "skipped": 0}, "passou"),
        (1, {"total": 5, "failed": 1, "errors": 0, "skipped": 0}, "falhou"),
        (0, {"total": 5, "failed": 0, "errors": 0, "skipped": 5}, "pulada"),
        (5, None, "pulada"),
        (None, None, "erro"),
        (2, None, "erro"),
    ],
)
def test_status_da_suite(code, tests, status) -> None:
    assert run_all.status_from(code, tests)[0] == status


def test_resumo_de_uma_avaliacao_de_qualidade(ragas_file: Path) -> None:
    detail, tests = run_all.quality_detail(ragas_file)

    # 2 casos x 5 métricas; c2 tem answer_relevancy 0,0 (medida, reprovada).
    assert tests == {"total": 10, "failed": 1, "errors": 0, "skipped": 0}
    assert "10/10 notas medidas" in detail


EXECUTION = {
    "started_at": "2026-10-07T11:30:00-03:00",
    "duration_seconds": 3000,
    "shared_answers": {"file": "respostas.json", "n": 10, "failures": 0, "generation_model": "gemma3:4b"},
    "reused_results": [],
    "steps": [
        {"key": "api", "name": "API (pytest)", "status": "passou", "detail": "22 passaram",
         "tests": {"total": 22, "failed": 0, "errors": 0, "skipped": 0}, "seconds": 3.0, "log": "logs/api.log"},
        {"key": "ragas_controles", "name": "RAGAS: controles do juiz", "status": "falhou", "detail": "1 falhou",
         "tests": {"total": 5, "failed": 1, "errors": 0, "skipped": 0}, "seconds": 60.0,
         "log": "logs/ragas_controles.log"},
        {"key": "ragas", "name": "RAGAS: avaliação", "status": "concluída", "detail": "50/50 notas medidas",
         "tests": {"total": 50, "failed": 30, "errors": 0, "skipped": 0}, "seconds": 900.0, "log": "logs/ragas.log"},
    ],
}


def test_relatorio_mostra_a_execucao_e_avisa_quando_os_controles_reprovam(ragas_file: Path, tmp_path: Path) -> None:
    outputs = build_report.build([ragas_file], ["md", "xlsx"], tmp_path / "out", execution=EXECUTION)

    md = outputs["md"].read_text(encoding="utf-8")
    assert "## 1. Execução dos testes" in md
    assert "✗ falhou" in md
    assert "✓ concluída" in md
    assert "mesmas respostas" in md
    assert "Controles do juiz RAGAS reprovaram" in md
    assert "Execução" in load_workbook(outputs["xlsx"]).sheetnames


def test_historico_guarda_cada_relatorio_e_conta_cada_rodada_uma_vez(
    ragas_file: Path, deepeval_file: Path, tmp_path: Path
) -> None:
    root = tmp_path / "out"
    build_report.build([ragas_file, deepeval_file], ["md"], root / "20261007T100000", history_root=root)
    # Segundo relatório a partir dos MESMOS resultados: nova entrada no índice,
    # mas as rodadas avaliadas não se duplicam na evolução.
    build_report.build([ragas_file, deepeval_file], ["md", "pdf"], root / "20261007T110000", history_root=root)

    entries = report_history.load_entries(root)
    assert [e["_folder"] for e in entries] == ["20261007T100000", "20261007T110000"]
    assert len(report_history.unique_runs(entries)) == 2

    index = (root / "HISTORICO.md").read_text(encoding="utf-8")
    assert index.index("20261007T110000") < index.index("20261007T100000")  # mais recente primeiro
    assert (root / "historico.csv").read_text(encoding="utf-8-sig").count("\n") == 1 + 2 * 5

    latest = root / report_history.LATEST_DIR
    assert (latest / "relatorio.pdf").is_file()
    assert json.loads((latest / "resumo.json").read_text(encoding="utf-8"))["files"]["pdf"] == "relatorio.pdf"


def test_relatorio_fora_da_pasta_de_historico_nao_mexe_no_indice(ragas_file: Path, tmp_path: Path) -> None:
    build_report.build([ragas_file], ["md"], tmp_path / "avulso")

    assert (tmp_path / "avulso" / "resumo.json").is_file()
    assert not (tmp_path / "HISTORICO.md").exists()
